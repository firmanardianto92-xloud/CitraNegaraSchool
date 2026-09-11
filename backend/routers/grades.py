"""Rapor nilai: guru input nilai per mata pelajaran, sistem hitung nilai akhir & rata-rata."""

import io
from typing import List, Optional

from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse

from lib.db import db
from models.schemas import (
    BatchGradeCreate, Grade, GradeCreate, GradeUpdate, KelasStudentRow, KelasSummary, LembarRapor,
    MapelRow, RaporSummary, SemesterPoint, TrenRapor,
)

router = APIRouter(prefix="/grades", tags=["grades"])

KKM = 75  # nilai di bawah ini dihitung sebagai "nilai merah"
CURRENT_SEMESTER = "Semester 1 - 2025/2026"
SEMESTER_ORDER = [
    "Semester 1 - 2024/2025",
    "Semester 2 - 2024/2025",
    "Semester 1 - 2025/2026",
]

# Mata pelajaran bawaan per jurusan/jenjang agar guru tidak perlu mengetik nama mapel.
MAPEL_UMUM = ["Matematika", "PPKN", "Bahasa Indonesia", "Bahasa Inggris", "Kesenian"]
MAPEL_BY_JURUSAN = {
    "IPA": MAPEL_UMUM + ["Fisika", "Kimia", "Biologi"],
    "IPS": MAPEL_UMUM + ["Ekonomi", "Sosiologi", "Geografi", "Sejarah"],
    "Bahasa": MAPEL_UMUM + ["Sastra Indonesia", "Bahasa Jepang", "Bahasa Arab", "Antropologi"],
    "RPL": MAPEL_UMUM + ["Pemrograman Dasar", "Basis Data", "Pemrograman Web", "Produk Kreatif"],
    "TKJ": MAPEL_UMUM + ["Jaringan Komputer", "Sistem Operasi", "Administrasi Server"],
    "Multimedia": MAPEL_UMUM + ["Desain Grafis", "Videografi", "Animasi 2D"],
    "Akuntansi": MAPEL_UMUM + ["Akuntansi Dasar", "Perpajakan", "Spreadsheet"],
}
MAPEL_BY_LEVEL = {
    "sd": MAPEL_UMUM + ["IPAS", "Pendidikan Agama", "Penjaskes"],
    "smp": MAPEL_UMUM + ["IPA Terpadu", "IPS Terpadu", "Prakarya", "Penjaskes"],
}


def _mapel_untuk(siswa: dict) -> List[str]:
    jurusan = (siswa or {}).get("jurusan") or ""
    if jurusan in MAPEL_BY_JURUSAN:
        return MAPEL_BY_JURUSAN[jurusan]
    return MAPEL_BY_LEVEL.get((siswa or {}).get("level") or "", MAPEL_UMUM)


def _sem_key(name: str) -> int:
    return SEMESTER_ORDER.index(name) if name in SEMESTER_ORDER else 99


def _predikat(nilai: float) -> str:
    if nilai >= 90:
        return "A"
    if nilai >= 80:
        return "B"
    if nilai >= 70:
        return "C"
    if nilai >= 60:
        return "D"
    return "E"


def _hitung(nilai_tugas: float, nilai_ujian: float) -> tuple[float, str]:
    akhir = round(nilai_tugas * 0.4 + nilai_ujian * 0.6, 1)
    return akhir, _predikat(akhir)


@router.get("", response_model=List[Grade])
async def list_grades(student_id: Optional[str] = None, semester: Optional[str] = None):
    query = {}
    if student_id:
        query["student_id"] = student_id
    if semester:
        query["semester"] = semester
    rows = await db.grades.find(query).sort("created_at", 1).to_list(500)
    return [Grade(**r) for r in rows]


@router.get("/rapor/{student_id}", response_model=RaporSummary)
async def rapor(student_id: str, semester: Optional[str] = None):
    rows = await db.grades.find(
        {"student_id": student_id, "semester": semester or CURRENT_SEMESTER}
    ).sort("created_at", 1).to_list(200)
    if not rows:
        siswa = await db.users.find_one({"id": student_id})
        return RaporSummary(
            student_id=student_id,
            student_name=siswa["name"] if siswa else "Siswa",
            semester=semester or CURRENT_SEMESTER,
            rata_rata=0,
            predikat="-",
            jumlah_mapel=0,
            grades=[],
        )
    grades = [Grade(**r) for r in rows]
    rata = round(sum(g.nilai_akhir for g in grades) / len(grades), 1)
    return RaporSummary(
        student_id=student_id,
        student_name=grades[0].student_name,
        semester=grades[0].semester,
        rata_rata=rata,
        predikat=_predikat(rata),
        jumlah_mapel=len(grades),
        grades=grades,
    )


@router.get("/kelas/{teacher_id}", response_model=KelasSummary)
async def kelas(teacher_id: str):
    """Rekap nilai seluruh siswa yang diampu seorang guru, termasuk penanda nilai merah."""
    guru = await db.users.find_one({"id": teacher_id, "role": {"$in": ["guru", "admin"]}})
    if not guru:
        raise HTTPException(status_code=404, detail="Guru tidak ditemukan")

    siswa = await db.users.find({"role": "siswa", "homeroom_teacher_id": teacher_id}).sort("name", 1).to_list(200)
    rows: List[KelasStudentRow] = []
    for s in siswa:
        g_rows = await db.grades.find(
            {"student_id": s["id"], "semester": CURRENT_SEMESTER}
        ).sort("subject", 1).to_list(100)
        grades = [Grade(**g) for g in g_rows]
        merah = [g for g in grades if g.nilai_akhir < KKM]
        rata = round(sum(g.nilai_akhir for g in grades) / len(grades), 1) if grades else 0
        rows.append(KelasStudentRow(
            student_id=s["id"],
            student_name=s["name"],
            kelas=s.get("kelas") or "",
            jurusan=s.get("jurusan") or "",
            rata_rata=rata,
            predikat=_predikat(rata) if grades else "-",
            jumlah_mapel=len(grades),
            jumlah_merah=len(merah),
            tuntas=len(merah) == 0,
            grades=grades,
        ))

    rata_kelas = round(sum(r.rata_rata for r in rows) / len(rows), 1) if rows else 0
    return KelasSummary(
        teacher_id=teacher_id,
        teacher_name=guru["name"],
        kelas=guru.get("kelas") or "",
        jurusan=guru.get("jurusan") or "",
        kkm=KKM,
        jumlah_siswa=len(rows),
        rata_rata_kelas=rata_kelas,
        siswa_bermasalah=len([r for r in rows if not r.tuntas]),
        students=rows,
    )


@router.get("/tren/{student_id}", response_model=TrenRapor)
async def tren(student_id: str):
    """Tren nilai antar semester untuk grafik perkembangan siswa."""
    siswa = await db.users.find_one({"id": student_id})
    rows = await db.grades.find({"student_id": student_id}).to_list(500)
    if not rows:
        return TrenRapor(
            student_id=student_id,
            student_name=(siswa or {}).get("name", "Siswa"),
            kkm=KKM,
        )

    grades = [Grade(**r) for r in rows]
    per_sem: dict[str, list[Grade]] = {}
    for g in grades:
        per_sem.setdefault(g.semester, []).append(g)

    semesters: List[SemesterPoint] = []
    mapel_series: List[dict] = []
    for sem in sorted(per_sem.keys(), key=_sem_key):
        items = per_sem[sem]
        rata = round(sum(i.nilai_akhir for i in items) / len(items), 1)
        semesters.append(SemesterPoint(
            semester=sem,
            rata_rata=rata,
            predikat=_predikat(rata),
            jumlah_mapel=len(items),
            jumlah_merah=len([i for i in items if i.nilai_akhir < KKM]),
        ))
        point = {"semester": sem.replace("Semester ", "Sem ").replace(" - ", " "), "Rata-rata": rata}
        for i in items:
            point[i.subject] = i.nilai_akhir
        mapel_series.append(point)

    selisih = round(semesters[-1].rata_rata - semesters[-2].rata_rata, 1) if len(semesters) > 1 else 0
    return TrenRapor(
        student_id=student_id,
        student_name=grades[0].student_name,
        kkm=KKM,
        semesters=semesters,
        mapel_series=mapel_series,
        selisih_terakhir=selisih,
    )


@router.get("/kelas/{teacher_id}/export")
async def export_kelas(teacher_id: str):
    """Unduh rekap nilai kelas sebagai file Excel untuk rapat wali murid."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    data = await kelas(teacher_id)

    wb = Workbook()
    ws = wb.active
    ws.title = "Rekap Nilai"

    ws["A1"] = "REKAP NILAI KELAS - SEKOLAH CITRA NEGARA"
    ws["A1"].font = Font(size=14, bold=True)
    ws["A2"] = f"{data.teacher_name} · {data.kelas} · Jurusan {data.jurusan or '-'} · KKM {data.kkm}"
    ws["A3"] = (f"Jumlah siswa: {data.jumlah_siswa} · Rata-rata kelas: {data.rata_rata_kelas} · "
                f"Siswa dengan nilai merah: {data.siswa_bermasalah}")

    mapel = sorted({g.subject for s in data.students for g in s.grades})
    header = ["No", "Nama Siswa", "Kelas"] + mapel + ["Rata-rata", "Predikat", "Nilai Merah", "Status"]
    ws.append([])
    ws.append(header)

    head_fill = PatternFill("solid", fgColor="1E3A8A")
    for cell in ws[5]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    merah_fill = PatternFill("solid", fgColor="FEE2E2")
    for idx, s in enumerate(data.students, start=1):
        by_subject = {g.subject: g.nilai_akhir for g in s.grades}
        ws.append([idx, s.student_name, s.kelas] + [by_subject.get(m, "") for m in mapel] +
                  [s.rata_rata, s.predikat, s.jumlah_merah, "Tuntas" if s.tuntas else "Perlu perhatian"])
        row = ws.max_row
        for col, m in enumerate(mapel, start=4):
            val = by_subject.get(m)
            if isinstance(val, (int, float)) and val < data.kkm:
                ws.cell(row=row, column=col).fill = merah_fill

    widths = [5, 26, 12] + [12] * len(mapel) + [11, 10, 12, 18]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=5, column=i).column_letter].width = w
    ws.freeze_panes = "D6"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"rekap-nilai-{(data.kelas or 'kelas').replace(' ', '-').lower()}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/lembar/{student_id}", response_model=LembarRapor)
async def lembar(student_id: str, semester: Optional[str] = None):
    """Lembar pengisian rapor: semua mapel siswa sudah terdaftar + nilai lama terisi otomatis."""
    siswa = await db.users.find_one({"id": student_id, "role": "siswa"})
    if not siswa:
        raise HTTPException(status_code=404, detail="Siswa tidak ditemukan")

    sem = semester or CURRENT_SEMESTER
    existing_rows = await db.grades.find({"student_id": student_id, "semester": sem}).to_list(100)
    existing = {r["subject"]: r for r in existing_rows}

    subjects = list(_mapel_untuk(siswa))
    # Mapel di luar daftar bawaan (mis. input manual sebelumnya) tetap ikut ditampilkan.
    for extra in existing:
        if extra not in subjects:
            subjects.append(extra)

    rows: List[MapelRow] = []
    for subject in subjects:
        g = existing.get(subject)
        rows.append(MapelRow(
            subject=subject,
            nilai_tugas=g["nilai_tugas"] if g else None,
            nilai_ujian=g["nilai_ujian"] if g else None,
            nilai_akhir=g["nilai_akhir"] if g else None,
            predikat=g["predikat"] if g else None,
            catatan=(g or {}).get("catatan", "") or "",
            grade_id=g["id"] if g else None,
            terisi=bool(g),
        ))

    return LembarRapor(
        student_id=student_id,
        student_name=siswa["name"],
        kelas=siswa.get("kelas") or "",
        jurusan=siswa.get("jurusan") or "",
        level=siswa.get("level") or "",
        semester=sem,
        kkm=KKM,
        jumlah_terisi=len([r for r in rows if r.terisi]),
        rows=rows,
    )


@router.post("/batch", response_model=List[Grade])
async def batch_grades(payload: BatchGradeCreate):
    """Simpan seluruh nilai satu siswa dalam sekali kirim (upsert per mata pelajaran)."""
    siswa = await db.users.find_one({"id": payload.student_id, "role": "siswa"})
    if not siswa:
        raise HTTPException(status_code=404, detail="Siswa tidak ditemukan")

    sem = payload.semester or CURRENT_SEMESTER
    tersimpan: List[Grade] = []
    for item in payload.items:
        if item.nilai_tugas is None and item.nilai_ujian is None:
            continue  # baris dibiarkan kosong oleh guru
        tugas = float(item.nilai_tugas or 0)
        ujian = float(item.nilai_ujian or 0)
        for label, val in (("tugas", tugas), ("ujian", ujian)):
            if val < 0 or val > 100:
                raise HTTPException(status_code=400, detail=f"Nilai {label} {item.subject} harus 0-100")

        akhir, pred = _hitung(tugas, ujian)
        existing = await db.grades.find_one(
            {"student_id": payload.student_id, "subject": item.subject, "semester": sem}
        )
        doc = Grade(
            id=existing["id"] if existing else Grade().id,
            student_id=payload.student_id,
            student_name=siswa["name"],
            subject=item.subject,
            semester=sem,
            nilai_tugas=tugas,
            nilai_ujian=ujian,
            nilai_akhir=akhir,
            predikat=pred,
            catatan=item.catatan or "",
            guru=payload.guru,
        )
        await db.grades.replace_one({"id": doc.id}, doc.model_dump(), upsert=True)
        tersimpan.append(doc)

    if not tersimpan:
        raise HTTPException(status_code=400, detail="Belum ada nilai yang diisi")
    return tersimpan


@router.post("", response_model=Grade)
async def create_grade(payload: GradeCreate):
    data = payload.model_dump()
    akhir, predikat = _hitung(data["nilai_tugas"], data["nilai_ujian"])
    grade = Grade(**data, nilai_akhir=akhir, predikat=predikat)
    await db.grades.insert_one(grade.model_dump())
    return grade


@router.patch("/{grade_id}", response_model=Grade)
async def update_grade(grade_id: str, payload: GradeUpdate):
    row = await db.grades.find_one({"id": grade_id})
    if not row:
        raise HTTPException(status_code=404, detail="Nilai tidak ditemukan")
    updates = {k: v for k, v in payload.model_dump().items() if v is not None}
    if not updates:
        raise HTTPException(status_code=400, detail="Tidak ada perubahan")
    merged = {**row, **updates}
    akhir, predikat = _hitung(merged["nilai_tugas"], merged["nilai_ujian"])
    updates.update({"nilai_akhir": akhir, "predikat": predikat})
    row = await db.grades.find_one_and_update({"id": grade_id}, {"$set": updates}, return_document=True)
    return Grade(**row)
